"""Read saved workbooks or take a temporary snapshot of desktop Excel."""

import io
import sys
from contextlib import contextmanager
from pathlib import Path
from tempfile import TemporaryDirectory

import openpyxl


SNAPSHOT_SHEET = 'ConverterInput'


def snapshot_directory():
    """Return the parent for the snapshot temp directory, or None for system temp.

    Mac Excel is sandboxed, so saving into /private/var/folders/.../T/ triggers a
    "파일 접근 권한 부여" (Grant File Access) prompt and fails with OSERROR -50 when
    it is cancelled. Its own container is writable without a prompt, so keep an
    own folder there instead of Data/tmp, which Office sweeps for its scratch files.
    """
    if sys.platform != 'darwin':
        return None
    container = Path.home() / 'Library/Containers/com.microsoft.Excel/Data'
    if not container.is_dir():
        return None
    directory = container / 'SpDocumentsConverter'
    try:
        directory.mkdir(exist_ok=True)
    except OSError:
        # An unwritable container or a stray file of that name is not worth an
        # error box; fall back to the system temp folder (which may prompt).
        return None
    return directory


def sheet_names(path):
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        return workbook.sheetnames
    finally:
        workbook.close()


@contextmanager
def open_sheet(path, name):
    if not path or not Path(path).is_file():
        raise ValueError('변환할 Excel 파일을 먼저 선택해주세요.')
    if not name:
        raise ValueError('변환할 시트를 선택해주세요.')
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        yield workbook[name]
    finally:
        workbook.close()


def active_book():
    # File-based conversion must not initialize Excel automation (or its exit hooks).
    import xlwings

    app = xlwings.apps.active
    book = app.books.active if app is not None else None
    if book is None:
        raise ValueError('Excel을 실행하고 변환할 문서를 열어주세요.')
    return book


@contextmanager
def active_sheet():
    book = active_book()
    sheet = book.sheets.active
    with TemporaryDirectory(prefix='spdocuments-', dir=snapshot_directory()) as directory:
        path = Path(directory) / 'input.xlsx'
        snapshot = None
        try:
            # Use the source book's Excel instance, including on Windows with
            # multiple instances. Always save as xlsx, even for unsaved/xls books.
            snapshot = book.app.books.add()
            sheet.copy(before=snapshot.sheets[0], name=SNAPSHOT_SHEET)
            # Only the disposable copy loses macros when saved as xlsx.
            with book.app.properties(display_alerts=False):
                snapshot.save(str(path))
        finally:
            try:
                if snapshot is not None:
                    snapshot.close()
            finally:
                book.activate()
        # Read the disposable snapshot through memory so a partially consumed
        # read-only iterator cannot retain a Windows handle to input.xlsx.
        with io.BytesIO(path.read_bytes()) as buffer:
            workbook = openpyxl.load_workbook(buffer, read_only=True, data_only=True)
            try:
                yield workbook[SNAPSHOT_SHEET]
            finally:
                workbook.close()


def selected_range():
    book = active_book()
    selection = book.app.selection
    if selection is None:
        raise ValueError('Excel에서 변환할 셀 영역을 선택해주세요.')
    data = [book.sheets.active.range('TITLES').value]
    data.extend(selection.value)
    return data
